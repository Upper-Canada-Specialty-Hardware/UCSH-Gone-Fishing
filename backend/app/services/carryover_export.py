"""The carry-over and payout export on the admin dashboard (#176).

Builds one Excel workbook HR can open, sort and send on: who has how many
vacation, carry-over and payout days, and which carry-over and payout requests
were made. Pure: takes the Staff Directory items and the carry-over/payout list
items already read, so it is tested without SharePoint or Postgres.
"""

from datetime import date, datetime
from io import BytesIO
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

TORONTO = ZoneInfo("America/Toronto")

BALANCE_HEADERS = [
    "Employee", "Email", "Department", "Location", "Employment type",
    "Vacation days left", "Carry Over days", "Payout days",
]
REQUEST_HEADERS = ["Request #", "Employee", "Type", "Days", "Status", "Submitted"]


def _num(value) -> float:
    """A balance as a number; blanks and junk count as 0."""
    try:
        return float(value or 0)
    except (TypeError, ValueError):
        return 0.0


def _submitted_on(fields: dict) -> date | None:
    """The Toronto date a request was made, from its Created time.

    Args:
        fields: The request item's fields.

    Returns:
        The date, or None when Created is missing or unreadable.
    """
    raw = fields.get("Created") or ""
    try:
        moment = datetime.fromisoformat(str(raw).replace("Z", "+00:00"))
    except ValueError:
        return None
    if moment.tzinfo is None:
        moment = moment.replace(tzinfo=ZoneInfo("UTC"))         # stored times are UTC
    return moment.astimezone(TORONTO).date()


def _autosize(sheet) -> None:
    """Widen each column to its longest value, within reason."""
    for i, column in enumerate(sheet.columns, start=1):
        longest = max((len(str(c.value)) for c in column if c.value is not None), default=8)
        sheet.column_dimensions[get_column_letter(i)].width = min(max(longest + 2, 10), 45)


def build_workbook(staff: list[dict], requests: list[dict], window: tuple[date, date] | None) -> bytes:
    """Build the export as .xlsx bytes.

    Args:
        staff: Staff Directory items ({"id", "fields"}).
        requests: Carry-over/payout list items ({"id", "fields"}).
        window: (start, end) set by the admin; when given, only requests made
            inside it are listed. None lists every request.

    Returns:
        The workbook file's bytes.
    """
    names = {str(s.get("id")): (s.get("fields", {}).get("Title") or "") for s in staff}
    book = Workbook()

    balances = book.active
    balances.title = "Balances"
    balances.append(BALANCE_HEADERS)
    for item in sorted(staff, key=lambda s: (s.get("fields", {}).get("Title") or "").lower()):
        f = item.get("fields", {})
        balances.append([
            f.get("Title", ""), f.get("EmailAddress", ""), f.get("Department", ""),
            f.get("Location", ""), f.get("SalaryHourly", ""),
            _num(f.get("CurrentVacationBalance")), _num(f.get("CarryOver")), _num(f.get("Payout")),
        ])

    listed = book.create_sheet("Requests")
    listed.append(REQUEST_HEADERS)
    rows = []
    for item in requests:
        f = item.get("fields", {})
        made = _submitted_on(f)
        if window and (made is None or not window[0] <= made <= window[1]):
            continue                                           # outside the set dates
        rows.append([
            int(item["id"]) if str(item.get("id", "")).isdigit() else item.get("id"),
            names.get(str(f.get("EmployeeID", "")), ""),       # the request stores the staff record id
            f.get("TypeofRequest", ""), _num(f.get("Days")), f.get("Status", ""),
            made.isoformat() if made else "",
        ])
    for row in sorted(rows, key=lambda r: r[5], reverse=True):  # newest first
        listed.append(row)

    for sheet in (balances, listed):
        for cell in sheet[1]:
            cell.font = Font(bold=True)                        # header row
        sheet.freeze_panes = "A2"                              # header stays put when scrolling
        _autosize(sheet)

    out = BytesIO()
    book.save(out)
    return out.getvalue()


def export_filename(today: date) -> str:
    """The download's file name, dated so repeated exports do not collide."""
    return f"carry-over-and-payout-{today.isoformat()}.xlsx"
