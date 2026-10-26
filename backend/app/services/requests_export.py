"""The requests Excel export (#278, rebuilt for #488).

The export dialog on the admin requests screen and the manager History tab
works out every cell itself, with the same words and dates the table shows
("Apr 21, 2026", "12 vacation", "Sick or personal"), and posts the sheets it
previewed. This module only writes them: a bold, frozen, filtered header, real
numbers in number columns with an optional totals row, wrapped long text, sized
columns, and no text that Excel could run as a formula. Pure: it reads no live
data, so it is tested without SharePoint or Postgres.
"""

import re
from datetime import date
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from pydantic import BaseModel, Field

# Characters that make Excel/Sheets read a cell as a live formula if they lead
# it. openpyxl writes a string starting with "=" as a real formula, so a field
# typed on a form (e.g. "=HYPERLINK(...)") would run when the export is opened.
_FORMULA_TRIGGERS = ("=", "+", "-", "@", "\t", "\r")

# Characters Excel refuses in a sheet name, and its name length limit.
_BAD_TITLE = re.compile(r"[\[\]:*?/\\]")
_TITLE_MAX = 31


class ExportColumn(BaseModel):
    """One column of a sheet, as the dialog chose it."""

    header: str = Field(max_length=80)
    number: bool = False                     # a real number: right-aligned, may be totalled
    total: bool = False                      # sum it in the totals row (number columns only)
    wrap: bool = False                       # long free text: wrapped and kept narrow


class ExportSheet(BaseModel):
    """One sheet: its title, columns and the rows the dialog previewed."""

    title: str = Field(max_length=80)
    columns: list[ExportColumn] = Field(min_length=1, max_length=60)
    rows: list[list[str | float | int | None]] = Field(max_length=50000)


class RequestsExportBody(BaseModel):
    """What the export dialog posts: the finished sheets, in order."""

    sheets: list[ExportSheet] = Field(min_length=1, max_length=10)


def _sanitize_cell_text(text: str) -> str:
    """Make a leading-formula string safe by keeping it as text (#294 S1).

    Args:
        text: A cell's string value.

    Returns:
        The text with a single quote in front when it starts with a formula
        trigger, which Excel treats as plain text; the text unchanged otherwise.
    """
    if text.startswith(_FORMULA_TRIGGERS):                     # would otherwise be read as a formula
        return "'" + text                                      # quote prefix: Excel keeps it as text
    return text


def _cell(value, number: bool):
    """One posted value as the cell to write.

    Args:
        value: Text, a number, or None.
        number: Whether the column holds numbers.

    Returns:
        A float for a readable number in a number column (so Excel sums it),
        None for a blank, otherwise sanitized text.
    """
    if value is None or value == "":
        return None                                            # an empty cell, not ""
    if number:
        try:
            return float(value)                                # Excel sums real numbers only
        except (TypeError, ValueError):
            pass                                               # not a number after all: keep as text
    return _sanitize_cell_text(str(value))                     # text never runs as a formula


def _safe_title(title: str, taken: set[str]) -> str:
    """A sheet name Excel accepts, unique within the workbook.

    Args:
        title: The name the dialog asked for.
        taken: Names already used (lower case; Excel ignores case).

    Returns:
        The name without forbidden characters, cut to 31, numbered if repeated.
    """
    base = _BAD_TITLE.sub(" ", title).strip()[:_TITLE_MAX] or "Requests"
    name, n = base, 2
    while name.lower() in taken:                               # Excel refuses two sheets of one name
        suffix = f" ({n})"
        name = base[:_TITLE_MAX - len(suffix)] + suffix
        n += 1
    taken.add(name.lower())
    return name


def _autosize(sheet, wrap_cols: set[int]) -> None:
    """Widen each column to its longest value, within reason.

    Args:
        sheet: The worksheet to size.
        wrap_cols: 1-based indexes of wrapped columns, capped narrower.
    """
    for i, column in enumerate(sheet.columns, start=1):
        longest = max((len(str(c.value)) for c in column if c.value is not None), default=8)
        cap = 40 if i in wrap_cols else 45                     # wrapped columns stay narrower
        sheet.column_dimensions[get_column_letter(i)].width = min(max(longest + 2, 10), cap)


def build_requests_workbook(sheets: list[ExportSheet]) -> bytes:
    """Write the dialog's sheets as .xlsx bytes.

    Args:
        sheets: The sheets in order, each with its columns and rows.

    Returns:
        The workbook file's bytes.
    """
    book = Workbook()
    book.remove(book.active)                                   # drop the default sheet; add our own
    taken: set[str] = set()
    bold = Font(bold=True)
    wrap = Alignment(wrap_text=True, vertical="top")
    for spec in sheets:
        sheet = book.create_sheet(_safe_title(spec.title, taken))
        cols = spec.columns
        width = len(cols)
        sheet.append([_sanitize_cell_text(c.header) for c in cols])   # header row
        for row in spec.rows:
            padded = (list(row) + [None] * width)[:width]      # exactly one value per column
            sheet.append([_cell(v, c.number) for v, c in zip(padded, cols)])
        data_rows = len(spec.rows)

        # A totals row under the columns the dialog asked to sum, when there are rows.
        totalled = [i for i, c in enumerate(cols) if c.number and c.total]
        if totalled and data_rows:
            total_row: list = [None] * width
            total_row[0] = "Total"
            for i in totalled:
                total_row[i] = sum(
                    v for v in (sheet.cell(row=r, column=i + 1).value for r in range(2, data_rows + 2))
                    if isinstance(v, (int, float))
                )
            sheet.append(total_row)
            for cell in sheet[sheet.max_row]:
                cell.font = bold                               # the totals row stands out

        for cell in sheet[1]:
            cell.font = bold                                   # header row
        sheet.freeze_panes = "A2"                              # header stays put when scrolling
        # Filter the header and the data only, so a sort never moves the totals row.
        sheet.auto_filter.ref = f"A1:{get_column_letter(width)}{data_rows + 1}"
        wrap_cols = {i + 1 for i, c in enumerate(cols) if c.wrap}
        for col_index in wrap_cols:                            # long free text wraps
            for (cell,) in sheet.iter_rows(min_col=col_index, max_col=col_index, min_row=2):
                cell.alignment = wrap
        for i, c in enumerate(cols, start=1):
            if c.number:                                       # numbers line up on the right
                for (cell,) in sheet.iter_rows(min_col=i, max_col=i, min_row=2):
                    cell.alignment = Alignment(horizontal="right")
        _autosize(sheet, wrap_cols)

    out = BytesIO()
    book.save(out)
    return out.getvalue()


def requests_export_filename(today: date) -> str:
    """The download's file name, dated so repeated exports do not collide."""
    return f"requests-{today.isoformat()}.xlsx"
