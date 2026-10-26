"""The staff export on the admin dashboard (#277).

Builds the Excel workbook HR previews and downloads from the staff table: one
"Staff" sheet of the people and columns HR chose, in the order the table showed
them, with optional "Summary by department" and "Change history" sheets. Pure:
it takes the Staff Directory items, the resolved manager names and the change
rows already read, so it is tested without SharePoint or Postgres.

The column set mirrors the person panel's export dialog. Numbers are written as
real numbers (so Excel sums and pivots them) and a below-zero balance is shown
in red, matching the carry-over export's conventions.
"""

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

from app.services.export_dates import readable_date, readable_datetime  # dates as the dashboard shows them (#490)

# Leading characters that make Excel treat a text cell as a formula on open; a
# Staff Directory value that begins with one of these is stored as literal text
# so a value like "=cmd" is shown verbatim and never run (formula injection).
_DANGEROUS_PREFIXES = ("=", "+", "-", "@", "\t", "\r", "\n")

# Sheet names, used by both the builder and the route/preview so they agree.
STAFF_SHEET = "Staff"
SUMMARY_SHEET = "Summary by department"
HISTORY_SHEET = "Change history"

# Friendly label per Staff Directory column for the change-history sheet, so a
# raw SharePoint column name never reaches HR. Mirrors staff_entitlements.
HISTORY_FIELD_LABELS = {
    "CurrentVacationBalance": "Vacation",
    "CurrentSickDayBalance": "Sick",
    "CurrentOvertimeBalance": "Make-up",
    "CarryOver": "Carry-over",
    "Payout": "Payout",
    "DefaultYearlyVacationDays": "Vacation entitlement",
    "SickDayEntitlement": "Sick entitlement",
    "Title": "Name",
    "EmailAddress": "Email",
    "CellNumber": "Cell",
    "Extension": "Extension",
    "Birthday": "Birthday",
    "Location": "Location",
    "Department": "Department",
    "SalaryHourly": "Salary or hourly",
    "Supervisor": "Supervisor (text)",
    "SystemCheck": "System check",
    "RequestAllowDate": "Can book vacation until",
    "Comments": "Comments",
    "TitleLink": "Person link",
    "SupervisorLink": "Supervisor link",
    "AllManagers": "Managers",
}


def _num(value) -> float:
    """A balance as a number; blanks and junk count as 0."""
    try:
        return float(value or 0)
    except (TypeError, ValueError):
        return 0.0


def _text(value) -> str:
    """Any stored value as plain text; None becomes a blank string."""
    return "" if value is None else str(value)


def _managers_text(row: dict) -> str:
    """The row's manager names joined for a single cell."""
    return ", ".join(row.get("managers") or [])


def _harden_against_formulas(cells) -> None:
    """Force any text cell that looks like a formula to be stored as plain text.

    ``append`` classifies a string beginning with ``=`` as a formula, so Excel
    would evaluate a comment such as ``=cmd`` on open. Reclassifying the cell as
    a string makes Excel show the text verbatim and never run it.

    Args:
        cells: The cells of one just-appended row.
    """
    for cell in cells:
        value = cell.value                               # the value append stored
        if isinstance(value, str) and value[:1] in _DANGEROUS_PREFIXES:
            cell.data_type = "s"                         # store/display as literal text


# One column definition per export column key. ``get`` takes a row
# ``{"id", "fields", "managers"}`` and returns the cell value. ``number`` marks
# the numeric columns (right-aligned, summed, red when below zero). ``group`` is
# the section the picker lists it under, and matches the dialog on the frontend.
COLUMNS: dict[str, dict] = {
    "id": {"header": "ID", "group": "Identity", "number": True,
           "get": lambda r: int(r["id"]) if str(r.get("id", "")).isdigit() else r.get("id", "")},
    "name": {"header": "Name", "group": "Identity",
             "get": lambda r: _text(r["fields"].get("Title"))},
    "email": {"header": "Email", "group": "Identity",
              "get": lambda r: _text(r["fields"].get("EmailAddress"))},
    "cell": {"header": "Cell", "group": "Contact",
             "get": lambda r: _text(r["fields"].get("CellNumber"))},
    "extension": {"header": "Extension", "group": "Contact",
                  "get": lambda r: _text(r["fields"].get("Extension"))},
    "birthday": {"header": "Birthday", "group": "Contact",
                 "get": lambda r: readable_date(r["fields"].get("Birthday"))},
    "location": {"header": "Location", "group": "Placement",
                 "get": lambda r: _text(r["fields"].get("Location"))},
    "department": {"header": "Department", "group": "Placement",
                   "get": lambda r: _text(r["fields"].get("Department"))},
    "salary_hourly": {"header": "Salary or hourly", "group": "Placement",
                      "get": lambda r: _text(r["fields"].get("SalaryHourly"))},
    "supervisor_text": {"header": "Supervisor (text)", "group": "Placement",
                        "get": lambda r: _text(r["fields"].get("Supervisor"))},
    "managers": {"header": "Managers", "group": "Placement", "get": _managers_text},
    "vacation_entitlement": {"header": "Vacation entitlement", "group": "Entitlements and balances",
                             "number": True, "get": lambda r: _num(r["fields"].get("DefaultYearlyVacationDays"))},
    "vacation": {"header": "Vacation", "group": "Entitlements and balances", "number": True,
                 "get": lambda r: _num(r["fields"].get("CurrentVacationBalance"))},
    "sick_entitlement": {"header": "Sick entitlement", "group": "Entitlements and balances",
                         "number": True, "get": lambda r: _num(r["fields"].get("SickDayEntitlement"))},
    "sick": {"header": "Sick", "group": "Entitlements and balances", "number": True,
             "get": lambda r: _num(r["fields"].get("CurrentSickDayBalance"))},
    "make_up_days": {"header": "Make-up (days)", "group": "Entitlements and balances", "number": True,
                     "get": lambda r: _num(r["fields"].get("CurrentOvertimeBalance"))},
    "make_up_hours": {"header": "Make-up (hours)", "group": "Entitlements and balances", "number": True,
                      "get": lambda r: _num(r["fields"].get("CurrentOvertimeBalance")) * 8},
    "carry_over": {"header": "Carry-over", "group": "Entitlements and balances", "number": True,
                   "get": lambda r: _num(r["fields"].get("CarryOver"))},
    "payout": {"header": "Payout", "group": "Entitlements and balances", "number": True,
               "get": lambda r: _num(r["fields"].get("Payout"))},
    "system_check": {"header": "System check", "group": "Status",
                     "get": lambda r: _text(r["fields"].get("SystemCheck"))},
    "request_allow_date": {"header": "Can book vacation until", "group": "Status",
                          "get": lambda r: readable_date(r["fields"].get("RequestAllowDate"))},
    "comments": {"header": "Comments", "group": "Record",
                 "get": lambda r: _text(r["fields"].get("Comments"))},
    "modified": {"header": "Last modified", "group": "Record",
                 "get": lambda r: readable_datetime(r["fields"].get("Modified"))},
}

# The order columns are written in, whatever order the picker sent them.
COLUMN_ORDER = list(COLUMNS.keys())

# "name" is always present so a row is never anonymous; the dialog locks it on.
LOCKED_COLUMNS = ["name"]


def resolve_columns(keys: list[str]) -> list[str]:
    """The valid, de-duplicated column keys in canonical order, name forced in.

    Args:
        keys: The column keys the picker sent (any order, maybe with unknowns).

    Returns:
        Known keys in ``COLUMN_ORDER`` order, with the locked columns included.
    """
    wanted = set(keys) | set(LOCKED_COLUMNS)              # name is never dropped
    return [k for k in COLUMN_ORDER if k in wanted and k in COLUMNS]


def _autosize(sheet, wrap_cols: set[int] | None = None) -> None:
    """Widen each column to its longest value, within reason.

    Args:
        sheet: The worksheet to size.
        wrap_cols: 1-based column indexes that wrap text, so they are capped
            narrower instead of stretched to a long free-text value.
    """
    wrap_cols = wrap_cols or set()
    for i, column in enumerate(sheet.columns, start=1):
        longest = max((len(str(c.value)) for c in column if c.value is not None), default=8)
        cap = 32 if i in wrap_cols else 45               # wrapped columns stay narrow
        sheet.column_dimensions[get_column_letter(i)].width = min(max(longest + 2, 10), cap)


def _history_value(value) -> object:
    """A history before/after value for a cell: numbers as numbers, else text."""
    if isinstance(value, bool):                          # a bool is not a figure here
        return str(value)
    if isinstance(value, (int, float)):
        return value                                     # kept numeric so Excel sums it
    return _text(value)


def build_workbook(
    staff: list[dict],
    columns: list[str],
    include_summary: bool,
    history: list[dict],
) -> bytes:
    """Build the staff export as .xlsx bytes.

    Args:
        staff: Rows to list, already filtered and ordered as the table showed
            them. Each is ``{"id", "fields", "managers"}`` where ``fields`` are
            the Staff Directory columns and ``managers`` the resolved names.
        columns: The column keys HR chose (order is normalised, name forced in).
        include_summary: Add the per-department summary sheet.
        history: Change rows for the "Change history" sheet, newest first; an
            empty list leaves that sheet out.

    Returns:
        The workbook file's bytes.
    """
    cols = resolve_columns(columns)                      # valid keys in canonical order
    defs = [COLUMNS[k] for k in cols]
    book = Workbook()

    # ---- Staff sheet ----
    sheet = book.active
    sheet.title = STAFF_SHEET
    sheet.append([d["header"] for d in defs])            # header row
    number_at = [i for i, d in enumerate(defs) if d.get("number")]  # columns to sum/align
    for row in staff:
        values = [d["get"](row) for d in defs]
        sheet.append(values)
        written = sheet[sheet.max_row]                   # style the row just added
        _harden_against_formulas(written)                # neutralise formula-like text
        for i in number_at:
            cell = written[i]
            if isinstance(cell.value, (int, float)) and cell.value < 0:
                cell.font = Font(color="C2410C")         # below zero in the error colour
    # A totals row under the balance columns (never the ID column).
    totals_cols = [i for i in number_at if cols[i] != "id"]
    staff_has_totals = bool(totals_cols and staff)       # remembered for the filter range
    if staff_has_totals:
        total_row: list = [""] * len(defs)
        total_row[0] = "Total"
        for i in totals_cols:
            total_row[i] = sum(_num(d_get) for d_get in (defs[i]["get"](r) for r in staff))
        sheet.append(total_row)
        for cell in sheet[sheet.max_row]:
            cell.font = Font(bold=True)                  # the totals row stands out

    # ---- Summary by department sheet ----
    if include_summary:
        summary = book.create_sheet(SUMMARY_SHEET)
        summary.append([
            "Department", "People", "Vacation total", "Vacation average",
            "Sick total", "Make-up total", "Carry-over total", "Below zero",
        ])
        by_dept: dict[str, list[dict]] = {}
        for row in staff:
            key = _text(row["fields"].get("Department")) or "Not set"
            by_dept.setdefault(key, []).append(row)
        for dept in sorted(by_dept):
            people = by_dept[dept]
            vac = sum(_num(p["fields"].get("CurrentVacationBalance")) for p in people)
            sick = sum(_num(p["fields"].get("CurrentSickDayBalance")) for p in people)
            make_up = sum(_num(p["fields"].get("CurrentOvertimeBalance")) for p in people)
            carry = sum(_num(p["fields"].get("CarryOver")) for p in people)
            below = sum(
                1 for p in people
                if any(_num(p["fields"].get(c)) < 0 for c in (
                    "CurrentVacationBalance", "CurrentSickDayBalance",
                    "CurrentOvertimeBalance", "CarryOver",
                ))
            )
            summary.append([
                dept, len(people), vac, round(vac / len(people), 2),
                sick, make_up, carry, below,
            ])
            _harden_against_formulas(summary[summary.max_row])  # dept names are user data

    # ---- Change history sheet ----
    if history:
        hist = book.create_sheet(HISTORY_SHEET)
        hist.append(["Date", "Name", "Field", "Before", "After", "Reason"])
        for row in history:
            hist.append([
                readable_date(row.get("changed_at")),                 # "Apr 21, 2026", Toronto day (#490)
                _text(row.get("employee_name")),
                HISTORY_FIELD_LABELS.get(row.get("field", ""), _text(row.get("field"))),
                _history_value(row.get("old_value")),
                _history_value(row.get("new_value")),
                _text(row.get("reason")),
            ])
            _harden_against_formulas(hist[hist.max_row])  # name/value/reason are user data

    # ---- Shared styling: bold frozen header, wrapped free-text, sized columns ----
    wrap_names = {"comments", "managers", "supervisor_text"}      # long free-text columns wrap
    staff_wrap = {i + 1 for i, k in enumerate(cols) if k in wrap_names}  # 1-based column indexes
    wrap_align = Alignment(wrap_text=True, vertical="top")       # the wrap style, shared
    for ws in book.worksheets:
        for cell in ws[1]:
            cell.font = Font(bold=True)                  # header row
        ws.freeze_panes = "A2"                           # header stays put when scrolling
        if ws.title == STAFF_SHEET:
            # Filter the header and the people only; the totals row stays out so a
            # sort or filter never reorders or hides it.
            last_col = get_column_letter(max(len(defs), 1))
            last_row = ws.max_row - (1 if staff_has_totals else 0)
            ws.auto_filter.ref = f"A1:{last_col}{max(last_row, 1)}"
            for col_index in staff_wrap:                 # wrap the long free-text columns
                for cell in ws.iter_rows(min_col=col_index, max_col=col_index):
                    cell[0].alignment = wrap_align
            _autosize(ws, staff_wrap)
        else:
            ws.auto_filter.ref = ws.dimensions           # Excel filters on every column
            if ws.title == HISTORY_SHEET:                # the Reason column (6th) wraps
                for cell in ws.iter_rows(min_col=6, max_col=6, min_row=2):
                    cell[0].alignment = wrap_align
            _autosize(ws)

    out = BytesIO()
    book.save(out)
    return out.getvalue()


def export_filename(today) -> str:
    """The download's file name, dated so repeated exports do not collide."""
    return f"staff-directory-{today.isoformat()}.xlsx"
