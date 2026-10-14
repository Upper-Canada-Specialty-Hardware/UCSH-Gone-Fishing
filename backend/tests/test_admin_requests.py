"""The admin requests screen's new, delete and export endpoints (#278).

The table, calendar and command bar are the frontend's; these cover the
backend they call: creating a request as an admin (Pending asks the manager,
Approved does not), deleting one with a reason and an optional refund while
keeping a record, the Excel export, and that an action on a request that no
longer exists is a clean 4xx rather than a 500.

Each router is mounted bare, without the admin-password guard, the way
test_admin_create_employee mounts the dashboard router.
"""

import asyncio

import httpx
import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient
from sqlalchemy import delete, select

from app.config import settings
from app.database import async_session
from app.models.request_deletion import RequestDeletion
from app.routes import admin_requests
from app.routes import dashboard
from app.services.requests_export import build_requests_workbook

# A fake Staff Directory record the endpoints resolve the employee to.
PAT = {"id": "7", "fields": {
    "Title": "Pat Staff", "EmailAddress": "pat@ucsh.com",
    "Department": "Warehouse", "Location": "Toronto Warden", "SalaryHourly": "Salary",
    "CurrentVacationBalance": 10, "CurrentSickDayBalance": 5, "CurrentOvertimeBalance": 2,
    "CarryOver": 0, "Payout": 0,
}}

VACATION_FORM = {"leave_type": "Vacation", "start_date": "2026-11-02", "end_date": "2026-11-03"}


def run(coro):
    """Run one coroutine to completion."""
    return asyncio.run(coro)


@pytest.fixture
def client(monkeypatch):
    """The admin-requests router alone, with SharePoint and the services stubbed.

    Attaches ``.created`` (submit_request calls) and ``.approved`` /
    ``.refunded`` (handler calls) so a test can see what the route did.
    """
    app = FastAPI()
    app.include_router(admin_requests.router)
    test_client = TestClient(app, raise_server_exceptions=False)
    test_client.created = []
    test_client.approved = []
    test_client.refunded = []
    test_client.recovered = []                                  # recovery (S3) calls, by item id
    test_client.claims = []                                     # (item id, action) pre-claims (S5)
    test_client.item = {"id": "500", "fields": {"Status": "Approved", "LeaveType": "Vacation", "Days": 2}}
    monkeypatch.setattr(settings, "PROCESSING_ENABLED", True)

    async def by_id(item_id):
        return PAT if str(item_id) == "7" else None

    async def submit(request_type, form, employee, source, notify=True):
        test_client.created.append({"type": request_type, "source": source, "notify": notify, "employee": employee["id"]})
        return {"id": "500"}

    async def approve_spy(request_id, approver):
        test_client.approved.append((request_id, approver))
        return {"status": "approved"}

    async def refund_spy(request_id, approver):
        test_client.refunded.append((request_id, approver))
        return {"status": "refunded"}

    async def recover_spy(request_id):
        test_client.recovered.append(str(request_id))
        return True                                             # handed to the manager as Pending

    async def claim_spy(list_id, item_id, action):
        test_client.claims.append((str(item_id), action))
        return True

    async def get_item(list_id, item_id):
        return test_client.item

    async def delete_item(list_id, item_id):
        test_client.item = None

    monkeypatch.setattr(admin_requests, "get_employee_by_id", by_id)
    monkeypatch.setattr(admin_requests, "submit_request", submit)
    monkeypatch.setattr(admin_requests, "claim_action", claim_spy)
    monkeypatch.setitem(admin_requests.APPROVE_HANDLERS, "leave", approve_spy)
    monkeypatch.setitem(admin_requests.REFUND_HANDLERS, "leave", refund_spy)
    monkeypatch.setitem(admin_requests.RECOVER_HANDLERS, "leave", recover_spy)
    monkeypatch.setattr(admin_requests.request_store, "get_list_item_or_none", get_item)
    monkeypatch.setattr(admin_requests.request_store, "delete_list_item", delete_item)
    return test_client


# ----- new request -----

def test_a_pending_request_is_created_and_the_manager_is_asked(client):
    resp = client.post("/admin/requests/new", json={"request_type": "leave", "employee_id": "7", "status": "Pending", "form": VACATION_FORM})
    assert resp.status_code == 200
    assert resp.json() == {"status": "submitted", "request_type": "leave", "item_id": "500"}
    # notify=True means it goes through the normal approval email, like any submission.
    assert client.created == [{"type": "leave", "source": "Admin dashboard", "notify": True, "employee": "7"}]
    assert client.approved == []                                # Pending is not approved here


def test_an_approved_request_is_approved_and_the_manager_is_not_asked(client):
    resp = client.post("/admin/requests/new", json={"request_type": "leave", "employee_id": "7", "status": "Approved", "form": VACATION_FORM})
    assert resp.status_code == 200
    assert resp.json() == {"status": "approved", "request_type": "leave", "item_id": "500"}
    # notify=False suppresses the manager's approval request; the admin approves it instead.
    assert client.created[0]["notify"] is False
    assert client.approved == [("500", "admin")]


def test_an_approval_that_is_refused_falls_back_to_pending(client, monkeypatch):
    # The admin's own approval was refused (a clash). Rather than leave the
    # request silently Pending with its manager email suppressed, hand it to the
    # manager the normal way and tell the admin (#294 S3). It is not an error.
    async def blocked(request_id, approver):
        return {"error": "Blocked by approved leave."}
    monkeypatch.setitem(admin_requests.APPROVE_HANDLERS, "leave", blocked)
    resp = client.post("/admin/requests/new", json={"request_type": "leave", "employee_id": "7", "status": "Approved", "form": VACATION_FORM})
    assert resp.status_code == 200
    body = resp.json()
    assert body["status"] == "pending" and body["item_id"] == "500"
    assert "manager" in body["message"]                          # the admin is told it went to the manager
    assert client.recovered == ["500"]                           # recovery ran: no orphaned, suppressed request
    assert client.approved == []                                 # the refused approval did not apply


def test_a_carryover_system_override_is_not_reported_as_approved(client, monkeypatch):
    # approve_carryover_payout can end a request by returning
    # {"status": "system_override_rejected"} (a terminal rejection), which must
    # not read as a success (#294 S2). Recovery returns False (it is terminal).
    async def overridden(request_id, approver):
        return {"status": "system_override_rejected"}

    async def recover_terminal(request_id):
        return False                                             # already rejected, not resurrected
    monkeypatch.setitem(admin_requests.APPROVE_HANDLERS, "carryover-payout", overridden)
    monkeypatch.setitem(admin_requests.RECOVER_HANDLERS, "carryover-payout", recover_terminal)
    form = {"type_of_request": "Payout", "days": 2}
    resp = client.post("/admin/requests/new", json={"request_type": "carryover-payout", "employee_id": "7", "status": "Approved", "form": form})
    assert resp.status_code == 200
    assert resp.json()["status"] == "pending"                    # specifically NOT "approved"


def test_an_approved_leave_with_no_days_is_not_approved(client):
    # process_new_leave_request logs and swallows a day-calculation failure, so a
    # day-costing leave could reach approval with 0 days and cost nothing. The
    # route re-reads the item and refuses that, falling into the S3 fallback (#294 S4).
    client.item = {"id": "500", "fields": {"Status": "Pending", "LeaveType": "Vacation", "Days": 0}}
    resp = client.post("/admin/requests/new", json={"request_type": "leave", "employee_id": "7", "status": "Approved", "form": VACATION_FORM})
    assert resp.status_code == 200
    body = resp.json()
    assert body["status"] == "pending"
    assert client.approved == []                                 # approval was never attempted
    assert client.recovered == ["500"]                           # handed to the manager instead


def test_a_bad_form_is_a_readable_400(client):
    bad = {**VACATION_FORM, "end_date": "2026-11-01"}           # end before start
    resp = client.post("/admin/requests/new", json={"request_type": "leave", "employee_id": "7", "status": "Pending", "form": bad})
    assert resp.status_code == 400 and "before" in resp.json()["detail"]
    assert client.created == []


def test_new_request_for_an_unknown_employee_is_a_404(client):
    resp = client.post("/admin/requests/new", json={"request_type": "leave", "employee_id": "999", "status": "Pending", "form": VACATION_FORM})
    assert resp.status_code == 404


def test_new_request_is_a_503_while_processing_is_off(client, monkeypatch):
    monkeypatch.setattr(settings, "PROCESSING_ENABLED", False)
    resp = client.post("/admin/requests/new", json={"request_type": "leave", "employee_id": "7", "status": "Pending", "form": VACATION_FORM})
    assert resp.status_code == 503
    assert client.created == []


# ----- delete -----

def _clear_deletions(list_item_id):
    async def _wipe():
        async with async_session() as session:
            await session.execute(delete(RequestDeletion).where(RequestDeletion.list_item_id == list_item_id))
            await session.commit()
    run(_wipe())


def _deletion(list_item_id):
    async def _get():
        async with async_session() as session:
            rows = (await session.execute(select(RequestDeletion).where(RequestDeletion.list_item_id == list_item_id))).scalars().all()
            return rows
    return run(_get())


def test_delete_needs_a_reason(client):
    resp = client.post("/admin/requests/leave/500/delete", json={"reason": "   ", "refund": False})
    assert resp.status_code == 400 and "reason" in resp.json()["detail"].lower()
    assert client.item is not None                              # nothing deleted


def test_delete_only_keeps_a_record_and_removes_the_item(client):
    _clear_deletions("500")
    resp = client.post("/admin/requests/leave/500/delete", json={"reason": "duplicate of 3905", "refund": False})
    assert resp.status_code == 200
    assert resp.json() == {"status": "deleted", "request_type": "leave", "request_id": "500", "refunded": False}
    assert client.item is None                                 # delete_list_item ran
    assert client.refunded == []                               # delete only, no refund
    rows = _deletion("500")
    assert len(rows) == 1 and rows[0].reason == "duplicate of 3905" and rows[0].refunded is False
    assert rows[0].fields.get("LeaveType") == "Vacation"       # the item's fields were kept
    _clear_deletions("500")


def test_refund_then_delete_refunds_first_and_records_it(client):
    _clear_deletions("500")
    resp = client.post("/admin/requests/leave/500/delete", json={"reason": "entered twice", "refund": True})
    assert resp.status_code == 200 and resp.json()["refunded"] is True
    assert client.refunded == [("500", "admin")]               # refund handler ran first
    assert client.item is None
    rows = _deletion("500")
    assert len(rows) == 1 and rows[0].refunded is True
    _clear_deletions("500")


def test_refund_is_skipped_for_a_request_that_was_not_approved(client):
    _clear_deletions("500")
    client.item = {"id": "500", "fields": {"Status": "Pending"}}   # nothing to refund
    resp = client.post("/admin/requests/leave/500/delete", json={"reason": "withdrawn", "refund": True})
    assert resp.status_code == 200 and resp.json()["refunded"] is False
    assert client.refunded == []
    _clear_deletions("500")


def test_deleting_a_pending_request_claims_approve_first(client):
    # A pending item could be approved by a manager mid-delete. Taking the
    # "approve" claim first stops that approval running against the item we are
    # about to delete (#294 S5).
    _clear_deletions("500")
    client.item = {"id": "500", "fields": {"Status": "Pending", "LeaveType": "Vacation"}}
    resp = client.post("/admin/requests/leave/500/delete", json={"reason": "withdrawn", "refund": False})
    assert resp.status_code == 200
    assert ("500", "approve") in client.claims                  # the in-flight approval is blocked
    _clear_deletions("500")


def test_a_refused_refund_stops_the_delete(client, monkeypatch):
    # The requested refund was refused, so nothing is deleted and no record is
    # kept: the request is left exactly as it was (#294 S5).
    _clear_deletions("500")

    async def refuse(request_id, approver):
        return {"error": "Not enough balance to refund."}
    monkeypatch.setitem(admin_requests.REFUND_HANDLERS, "leave", refuse)
    resp = client.post("/admin/requests/leave/500/delete", json={"reason": "oops", "refund": True})
    assert resp.status_code == 400 and "refund" in resp.json()["detail"].lower()
    assert client.item is not None                              # nothing deleted
    assert _deletion("500") == []                               # no record kept
    _clear_deletions("500")


def test_a_failed_delete_after_a_refund_is_a_502(client, monkeypatch):
    # The refund ran but the delete itself failed. Report a 502 and keep no
    # deletion record, so a record never outlives an item that is still there (#294 S5).
    _clear_deletions("500")

    async def boom(list_id, item_id):
        raise httpx.HTTPStatusError(
            "server error", request=httpx.Request("DELETE", "http://x"), response=httpx.Response(500),
        )
    monkeypatch.setattr(admin_requests.request_store, "delete_list_item", boom)
    resp = client.post("/admin/requests/leave/500/delete", json={"reason": "entered twice", "refund": True})
    assert resp.status_code == 502
    assert client.refunded == [("500", "admin")]               # the refund already ran
    assert _deletion("500") == []                               # no record, since the item is still there
    _clear_deletions("500")


def test_a_404_from_the_delete_is_treated_as_success(client, monkeypatch):
    # SharePoint answered 404: the item is already gone, which is the outcome we
    # wanted, so keep the record and report success (#294 S5).
    _clear_deletions("500")

    async def already_gone(list_id, item_id):
        raise httpx.HTTPStatusError(
            "not found", request=httpx.Request("DELETE", "http://x"), response=httpx.Response(404),
        )
    monkeypatch.setattr(admin_requests.request_store, "delete_list_item", already_gone)
    resp = client.post("/admin/requests/leave/500/delete", json={"reason": "already removed", "refund": False})
    assert resp.status_code == 200 and resp.json()["status"] == "deleted"
    assert len(_deletion("500")) == 1                           # the record is still kept
    _clear_deletions("500")


def test_delete_of_a_missing_request_is_a_404(client, monkeypatch):
    async def gone(list_id, item_id):
        return None
    monkeypatch.setattr(admin_requests.request_store, "get_list_item_or_none", gone)
    resp = client.post("/admin/requests/leave/500/delete", json={"reason": "x", "refund": False})
    assert resp.status_code == 404


def test_delete_is_a_503_while_processing_is_off(client, monkeypatch):
    monkeypatch.setattr(settings, "PROCESSING_ENABLED", False)
    resp = client.post("/admin/requests/leave/500/delete", json={"reason": "x", "refund": False})
    assert resp.status_code == 503
    assert client.item is not None


# ----- export -----

def test_the_export_has_one_sheet_per_list_with_every_column():
    from openpyxl import load_workbook
    from io import BytesIO

    requests = [
        {"request_type": "leave", "id": "1", "employee_name": "Avery", "managers": "Sam",
         "Status": "Approved", "Created": "2026-10-01", "ApprovedDate": "2026-10-02",
         "LeaveType": "Vacation", "StartDate": "2026-11-02", "EndDate": "2026-11-06", "Days": 5, "Title": "trip",
         "OddColumn": "kept"},
        {"request_type": "overtime", "id": "2", "employee_name": "Jamie", "managers": "Lee",
         "Status": "Pending", "Created": "2026-10-03", "StartDate": "2026-10-03", "Hours": 4, "Title": "inventory"},
        {"request_type": "carryover-payout", "id": "3", "employee_name": "Morgan", "managers": "Pat",
         "Status": "Approved", "Created": "2026-09-01", "TypeofRequest": "Payout", "Days": 5, "NewBalance": "x"},
    ]
    book = load_workbook(BytesIO(build_requests_workbook(requests)))
    assert book.sheetnames == ["Leave Requests", "Overtime Requests Data", "CarryOver_Payout_Data"]
    leave_headers = [c.value for c in book["Leave Requests"][1]]
    assert "Employee" in leave_headers and "Leave type" in leave_headers
    assert "OddColumn" in leave_headers                        # an unexpected field is still a column
    # The one leave row's employee was written.
    assert book["Leave Requests"]["B2"].value == "Avery"


def test_the_export_opens_even_with_no_requests():
    from openpyxl import load_workbook
    from io import BytesIO
    book = load_workbook(BytesIO(build_requests_workbook([])))
    assert book.sheetnames                                     # at least one sheet, so the file opens


def test_the_export_neutralises_a_formula_in_a_field():
    # A field typed on a form could start with "=", "+", "-" or "@", which Excel
    # runs as a live formula. The export prefixes a quote so it stays text (#294 S1).
    from openpyxl import load_workbook
    from io import BytesIO
    requests = [
        {"request_type": "leave", "id": "1", "employee_name": "Avery", "managers": "Sam",
         "Status": "Pending", "LeaveType": "Vacation", "Title": '=HYPERLINK("http://evil","click")'},
        {"request_type": "leave", "id": "2", "employee_name": "Blake", "managers": "Sam",
         "Status": "Pending", "LeaveType": "Vacation", "Title": "@SUM(A1:A9)"},
    ]
    sheet = load_workbook(BytesIO(build_requests_workbook(requests)))["Leave Requests"]
    headers = [c.value for c in sheet[1]]
    title_col = headers.index("Title") + 1                      # 1-based column of the Title field
    assert sheet.cell(row=2, column=title_col).value.startswith("'=")   # quoted, so inert
    assert sheet.cell(row=3, column=title_col).value.startswith("'@")


# ----- a missing request through the dashboard actions is a clean 4xx -----

@pytest.fixture
def dash_client(monkeypatch):
    """The dashboard router alone, for the admin approve/reject/remind endpoints."""
    monkeypatch.setattr(settings, "PROCESSING_ENABLED", True)
    app = FastAPI()
    app.include_router(dashboard.router)
    return TestClient(app, raise_server_exceptions=False)


def test_approving_a_deleted_request_is_a_404_not_a_500(dash_client, monkeypatch):
    # Approving by an old link after the item was deleted: the handler reads the
    # item and SharePoint answers 404; the route must turn that into a 404.
    async def gone(request_id, manager_id):
        raise httpx.HTTPStatusError("not found", request=httpx.Request("GET", "http://x"),
                                    response=httpx.Response(404))
    monkeypatch.setitem(dashboard.HANDLERS, ("leave", "approve"), gone)
    resp = dash_client.post("/admin/approve/leave/999")
    assert resp.status_code == 404 and "no longer exists" in resp.json()["detail"]


def test_reminding_a_deleted_request_is_a_404_not_a_500(dash_client, monkeypatch):
    async def gone(request_type, request_id):
        raise httpx.HTTPStatusError("not found", request=httpx.Request("GET", "http://x"),
                                    response=httpx.Response(404))
    monkeypatch.setattr("app.tasks.reminders.send_reminder_now", gone)
    resp = dash_client.post("/admin/send-reminder/leave/999")
    assert resp.status_code == 404
