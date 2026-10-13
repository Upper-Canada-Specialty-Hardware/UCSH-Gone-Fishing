"""The carry-over and payout Excel export (#176)."""

from datetime import date
from io import BytesIO

from openpyxl import load_workbook

from app.services.carryover_export import build_workbook, export_filename

STAFF = [
    {"id": "7", "fields": {"Title": "Pat Staff", "EmailAddress": "pat@ucsh.com", "Department": "Warehouse",
                           "Location": "Barrie", "SalaryHourly": "Salary",
                           "CurrentVacationBalance": 6, "CarryOver": 2, "Payout": 1.5}},
    {"id": "8", "fields": {"Title": "Ali Other", "CurrentVacationBalance": None}},
]
REQUESTS = [
    {"id": "301", "fields": {"EmployeeID": 7, "TypeofRequest": "Payout", "Days": 2, "Status": "Approved",
                             "Created": "2026-11-05T15:00:00Z"}},
    {"id": "302", "fields": {"EmployeeID": 7, "TypeofRequest": "Carry Over", "Days": 3, "Status": "Pending",
                             "Created": "2026-06-01T15:00:00Z"}},
]


def _open(data: bytes):
    """Read the workbook back."""
    return load_workbook(BytesIO(data))


def test_balances_list_everyone_sorted_with_blanks_as_zero():
    sheet = _open(build_workbook(STAFF, REQUESTS, None))["Balances"]
    rows = list(sheet.iter_rows(values_only=True))
    assert rows[0][0] == "Employee"
    assert rows[1][0] == "Ali Other" and rows[1][5] == 0               # blank balance counts as 0
    assert rows[2] == ("Pat Staff", "pat@ucsh.com", "Warehouse", "Barrie", "Salary", 6, 2, 1.5)


def test_requests_name_the_employee_and_keep_only_the_window():
    every = list(_open(build_workbook(STAFF, REQUESTS, None))["Requests"].iter_rows(values_only=True))
    assert len(every) == 3                                             # header + both requests
    window = (date(2026, 11, 1), date(2026, 11, 30))
    inside = list(_open(build_workbook(STAFF, REQUESTS, window))["Requests"].iter_rows(values_only=True))
    assert inside[1:] == [(301, "Pat Staff", "Payout", 2, "Approved", "2026-11-05")]


def test_the_file_name_is_dated():
    assert export_filename(date(2026, 11, 5)) == "carry-over-and-payout-2026-11-05.xlsx"
