"""Tests for the staff export workbook builder (#277).

Pure: build_workbook takes already-read Staff Directory rows, manager names and
change rows, so it needs no SharePoint or database. Each test reads the bytes
back with openpyxl and asserts the sheets, headers, values and styling.
"""
from datetime import date
from io import BytesIO

import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.routes.dashboard import router as dashboard_router
from app.services import staff_export as sx


def _rows() -> list[dict]:
    """Two people as the route hands them to the builder."""
    return [
        {
            "id": "601",
            "fields": {
                "Title": "Avery Example", "EmailAddress": "avery@ucsh.com",
                "Department": "Project Management", "Location": "Toronto Warden",
                "SalaryHourly": "Salary", "CurrentVacationBalance": 4.5,
                "CurrentSickDayBalance": 2, "CurrentOvertimeBalance": 0.25,
                "CarryOver": 0, "Payout": 0, "Comments": "a long note that should wrap",
            },
            "managers": ["Sam Supervisor"],
        },
        {
            "id": "602",
            "fields": {
                "Title": "Jamie Sample", "EmailAddress": "jamie@ucsh.com",
                "Department": "Warehouse", "Location": "British Columbia",
                "SalaryHourly": "Hourly BC", "CurrentVacationBalance": -1.5,
                "CurrentSickDayBalance": 4, "CurrentOvertimeBalance": 0,
                "CarryOver": 0, "Payout": 0,
            },
            "managers": [],
        },
    ]


def _open(data: bytes):
    """Load the produced workbook back for inspection."""
    return load_workbook(BytesIO(data))


def test_resolve_columns_forces_name_in_canonical_order():
    """Unknown keys drop out, name is always present, order is canonical."""
    cols = sx.resolve_columns(["carry_over", "email", "bogus"])
    assert "name" in cols                                  # name is locked on
    assert "bogus" not in cols                             # unknown dropped
    assert cols == [c for c in sx.COLUMN_ORDER if c in cols]  # canonical order


def test_staff_sheet_has_chosen_columns_and_values():
    """The Staff sheet lists the chosen columns with their values."""
    data = sx.build_workbook(_rows(), ["name", "email", "vacation"], False, [])
    ws = _open(data)[sx.STAFF_SHEET]
    assert [c.value for c in ws[1]] == ["Name", "Email", "Vacation"]  # header row
    assert ws["A2"].value == "Avery Example"
    assert ws["C2"].value == 4.5                           # stored as a number


def test_totals_row_sums_balances_and_marks_below_zero():
    """A totals row sums the balance columns and negatives are coloured."""
    data = sx.build_workbook(_rows(), ["name", "vacation"], False, [])
    ws = _open(data)[sx.STAFF_SHEET]
    # Rows: header(1), Avery(2), Jamie(3), Total(4).
    assert ws["A4"].value == "Total"
    assert ws["B4"].value == 3.0                           # 4.5 + (-1.5)
    assert ws["B3"].value == -1.5
    assert (ws["B3"].font.color.rgb or "").endswith("C2410C")  # below zero in error colour


def test_summary_sheet_groups_by_department():
    """The summary sheet has one row per department with its totals."""
    data = sx.build_workbook(_rows(), ["name"], True, [])
    ws = _open(data)[sx.SUMMARY_SHEET]
    headers = [c.value for c in ws[1]]
    assert headers[0] == "Department" and "Below zero" in headers
    depts = {ws.cell(row=r, column=1).value for r in range(2, ws.max_row + 1)}
    assert depts == {"Project Management", "Warehouse"}


def test_history_sheet_present_only_when_rows_given():
    """The change-history sheet appears only when history rows are passed."""
    none = _open(sx.build_workbook(_rows(), ["name"], False, []))
    assert sx.HISTORY_SHEET not in none.sheetnames
    history = [{
        "employee_name": "Avery Example", "field": "CurrentVacationBalance",
        "old_value": 10, "new_value": 4.5, "reason": "used days",
        "changed_at": "2026-10-14T13:00:00+00:00",
    }]
    ws = _open(sx.build_workbook(_rows(), ["name"], False, history))[sx.HISTORY_SHEET]
    assert [c.value for c in ws[1]] == ["Date", "Name", "Field", "Before", "After", "Reason"]
    assert ws["C2"].value == "Vacation"                    # column name -> friendly label
    assert ws["B2"].value == "Avery Example"
    assert ws["A2"].value == "Oct 14, 2026"                # the Toronto day, as the dashboard (#490)


def test_header_is_frozen_and_filtered():
    """Every sheet freezes and filters its header row."""
    ws = _open(sx.build_workbook(_rows(), ["name", "email"], False, []))[sx.STAFF_SHEET]
    assert ws.freeze_panes == "A2"
    assert ws.auto_filter.ref is not None                  # Excel filters enabled


def test_date_columns_are_written_as_the_dashboard_shows_them():
    """Birthday and the can-book date read as their own day, "Mar 3, 1990" (#490)."""
    rows = _rows()
    rows[0]["fields"]["Birthday"] = "1990-03-03T00:00:00Z"
    rows[0]["fields"]["RequestAllowDate"] = "2026-12-31T05:00:00Z"
    data = sx.build_workbook(rows, ["name", "birthday", "request_allow_date"], False, [])
    ws = _open(data)[sx.STAFF_SHEET]
    assert ws["B2"].value == "Mar 3, 1990"                 # a day saved at UTC midnight keeps its day
    assert ws["C2"].value == "Dec 31, 2026"                # Toronto midnight (05:00Z) is that day


def test_formula_like_text_is_stored_as_plain_text():
    """A comment that looks like a formula is stored as text, so Excel never runs it."""
    rows = _rows()
    rows[0]["fields"]["Comments"] = "=1+1"                  # a formula-injection attempt
    data = sx.build_workbook(rows, ["name", "comments"], False, [])
    ws = _open(data)[sx.STAFF_SHEET]
    assert ws["B2"].value == "=1+1"                        # the text is kept verbatim
    assert ws["B2"].data_type == "s"                       # as a string, not a formula


def test_totals_row_is_outside_the_filter_range():
    """The totals row sits below the auto-filter so a sort never reorders it."""
    data = sx.build_workbook(_rows(), ["name", "vacation"], False, [])
    ws = _open(data)[sx.STAFF_SHEET]
    # Rows: header(1), Avery(2), Jamie(3), Total(4); the filter covers 1..3 only.
    assert ws.auto_filter.ref == "A1:B3"


def test_wrap_columns_wrap_their_text():
    """The long free-text columns wrap so the sheet never has to scroll sideways."""
    data = sx.build_workbook(_rows(), ["name", "comments"], False, [])
    ws = _open(data)[sx.STAFF_SHEET]
    assert ws["B2"].alignment.wrap_text is True            # comments column wraps


def test_filename_is_dated():
    """The download name carries the date so repeated exports do not collide."""
    assert sx.export_filename(date(2026, 10, 14)) == "staff-directory-2026-10-14.xlsx"


# ---------------------------------------------------------------------------
# The export route: real route, the three reads it makes stubbed
# ---------------------------------------------------------------------------

@pytest.fixture
def export_client(monkeypatch):
    """The dashboard router with the export route's SharePoint reads stubbed.

    Returns:
        A TestClient whose ``/admin/staff/export`` reads two fixed people, an
        empty name map and no history, so only the route's ordering/selection
        logic is under test.
    """
    import app.routes.dashboard as dash
    from app.services import employee as emp_mod
    from app.services import staff_changes as sc_mod

    items = [
        {"id": "601", "fields": {"Title": "Avery Example", "Department": "PM"}},
        {"id": "602", "fields": {"Title": "Jamie Sample", "Department": "Warehouse"}},
    ]

    async def fake_items(list_id, *a, **k):
        return items

    async def fake_map():
        return {}

    async def fake_history(ids, limit=5000):
        return []

    monkeypatch.setattr(dash.sp_client, "get_list_items", fake_items)
    monkeypatch.setattr(emp_mod, "_get_sp_user_name_map", fake_map)
    monkeypatch.setattr(sc_mod, "list_changes_for", fake_history)
    app = FastAPI()
    app.include_router(dashboard_router)
    return TestClient(app, raise_server_exceptions=False)


def _staff_names(resp) -> list:
    """The Staff sheet's name column (data rows only) out of a response."""
    ws = load_workbook(BytesIO(resp.content))[sx.STAFF_SHEET]
    return [ws.cell(row=r, column=1).value for r in range(2, ws.max_row + 1)]


def test_export_route_null_ids_is_everyone_in_name_order(export_client):
    """Omitting employee_ids exports everyone, sorted by name."""
    resp = export_client.post("/admin/staff/export", json={"columns": ["name"]})
    assert resp.status_code == 200
    assert _staff_names(resp) == ["Avery Example", "Jamie Sample"]


def test_export_route_empty_ids_is_an_empty_sheet(export_client):
    """An explicit empty list means nobody, so only the header is written."""
    resp = export_client.post(
        "/admin/staff/export", json={"columns": ["name"], "employee_ids": []}
    )
    assert resp.status_code == 200
    assert _staff_names(resp) == []                         # header only, nobody listed


def test_export_route_keeps_order_and_drops_unknown_ids(export_client):
    """The sheet keeps the table's order and silently drops ids that do not exist."""
    resp = export_client.post(
        "/admin/staff/export",
        json={"columns": ["name"], "employee_ids": ["602", "999", "601"]},
    )
    assert resp.status_code == 200
    assert _staff_names(resp) == ["Jamie Sample", "Avery Example"]  # 999 dropped, order kept
