"""The manager History tab's team requests Excel export (#360).

A manager or admin can download the team requests on screen as an .xlsx; an
employee is refused. Driven through a TestClient so the real route body runs,
with the signed-token dependency overridden to a chosen role (the same pattern
as test_dashboard_postgres_routes.py).
"""

from io import BytesIO

from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.routes.dashboard import DashboardUser, get_dashboard_user
from app.routes.team_export import router as team_export_router

# The sheets the export dialog posts (#488): it works out each cell from the
# team rows on screen, so the route only writes them.
TEAM_SHEETS = [{
    "title": "Requests",
    "columns": [{"header": "Name"}, {"header": "Type"}, {"header": "Start date"},
                {"header": "Days", "number": True, "total": True}],
    "rows": [["Ana Reyes", "Vacation", "Nov 2, 2026", 5], ["Ben Okafor", "Make-up", "Oct 3, 2026", None]],
}]

XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _client(role: str) -> TestClient:
    """A TestClient for the export router alone, signed in as the given role.

    Args:
        role: "manager", "admin" or "employee" to stand in for the signed token.

    Returns:
        The client, with get_dashboard_user overridden to that role.
    """
    app = FastAPI()
    app.include_router(team_export_router)
    app.dependency_overrides[get_dashboard_user] = lambda: DashboardUser(role=role, user_id="42")
    return TestClient(app, raise_server_exceptions=False)


def test_a_manager_gets_an_xlsx():
    resp = _client("manager").post("/team/requests/export", json={"sheets": TEAM_SHEETS})
    assert resp.status_code == 200, resp.text
    assert resp.headers["content-type"] == XLSX_MEDIA_TYPE
    assert "attachment" in resp.headers["content-disposition"]
    # The bytes really are a workbook, with the team's rows as the dialog sent them.
    book = load_workbook(BytesIO(resp.content))
    assert book.sheetnames == ["Requests"]
    assert book["Requests"]["A2"].value == "Ana Reyes" and book["Requests"]["C2"].value == "Nov 2, 2026"


def test_an_admin_may_also_export():
    resp = _client("admin").post("/team/requests/export", json={"sheets": TEAM_SHEETS})
    assert resp.status_code == 200, resp.text
    assert resp.headers["content-type"] == XLSX_MEDIA_TYPE


def test_an_employee_is_refused():
    resp = _client("employee").post("/team/requests/export", json={"sheets": TEAM_SHEETS})
    assert resp.status_code == 403, resp.text
